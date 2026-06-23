import os
import openpyxl

def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    xlsx_path = os.path.abspath(os.path.join(script_dir, "..", "reference", "MahindraKycSheet.xlsx"))
    
    if not os.path.exists(xlsx_path):
        print(f"Excel file not found at: {xlsx_path}")
        return
        
    print(f"Loading workbook: {xlsx_path}")
    wb = openpyxl.load_workbook(xlsx_path)
    for name in wb.sheetnames:
        print(f"\nSheet: {name}")
        sheet = wb[name]
        for row in list(sheet.iter_rows(values_only=True))[:10]: # print first 10 rows
            print(row)

if __name__ == "__main__":
    main()
