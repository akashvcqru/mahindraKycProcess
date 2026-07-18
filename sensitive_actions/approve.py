import json
import os
import time
import logging
import openpyxl
import requests

def approve_claim(claim, webapp_url=None):
    """Mark a claim as APPROVED and sync changes to local JSON, local Excel, and Google Sheets."""
    logging.info(f"[Manual Approval] Approving claim for customer: {claim.get('customer_name')} (Row {claim.get('row_idx') + 1})")
    
    # 1. Update status in ui_history.json
    history_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "ui_history.json")
    if os.path.exists(history_path):
        try:
            with open(history_path, "r", encoding="utf-8") as f:
                history = json.load(f)
            for item in history:
                if item.get("row_idx") == claim.get("row_idx") and item.get("customer_name") == claim.get("customer_name"):
                    item["status"] = "APPROVED"
                    item["hold_reasons"] = ""
                    break
            with open(history_path, "w", encoding="utf-8") as f:
                json.dump(history, f, indent=2)
            logging.info("[Manual Approval] Updated ui_history.json successfully.")
        except Exception as e:
            logging.error(f"[Manual Approval] Failed to update history JSON: {e}")

    # 2. Update status in local Excel kyc_process_results.xlsx
    excel_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "kyc_process_results.xlsx")
    if os.path.exists(excel_path):
        try:
            wb = openpyxl.load_workbook(excel_path)
            ws = wb["KYC Results"]
            row_num = claim.get("row_idx") + 1
            updated = False
            for r in range(2, ws.max_row + 1):
                if ws.cell(row=r, column=1).value == row_num:
                    ws.cell(row=r, column=9, value="APPROVED")
                    ws.cell(row=r, column=10, value="")
                    updated = True
                    break
            if updated:
                wb.save(excel_path)
                logging.info("[Manual Approval] Updated local Excel spreadsheet successfully.")
        except Exception as e:
            logging.error(f"[Manual Approval] Failed to update local Excel: {e}")

    # 3. Post sync updates to Google Sheet Web App
    if webapp_url:
        record = {
            "row_num": claim.get("row_idx") + 1,
            "claim_no": claim.get("claim_details", {}).get("Invoice No") or claim.get("claim_details", {}).get("Invoice No.", ""),
            "claim_date": claim.get("claim_details", {}).get("Invoice Date", ""),
            "area_office": claim.get("claim_details", {}).get("Area Office", ""),
            "customer_name": claim.get("customer_name"),
            "dealer_name": claim.get("claim_details", {}).get("Dealer Name", ""),
            "dealer_branch": claim.get("claim_details", {}).get("Dealer Branch", ""),
            "scheme": claim.get("old_vehicle_details", {}).get("Scheme", ""),
            "status": "APPROVED",
            "hold_reasons": "",
            "date_processed": claim.get("date_processed") or time.strftime("%Y-%m-%d %H:%M:%S"),
            "chassis_no": claim.get("claim_details", {}).get("Chassis No") or claim.get("claim_details", {}).get("Chassis No.", ""),
            "duration": claim.get("duration", 0),
        }
        try:
            resp = requests.post(webapp_url, json=record, timeout=15)
            if resp.status_code == 200:
                logging.info("[Manual Approval] Google Sheet synchronized successfully.")
                return True
            else:
                logging.warning(f"[Manual Approval] Google Sheet sync returned HTTP {resp.status_code}")
        except Exception as e:
            logging.error(f"[Manual Approval] Google Sheet sync failed: {e}")
    return True
